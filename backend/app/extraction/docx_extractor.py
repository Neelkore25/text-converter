"""DOCX document extraction module (Phase 14).

Extracts text and paragraph structure from Microsoft Word (.docx) documents,
inspecting font hints from run properties.
"""
from __future__ import annotations

from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
from backend.app.extraction import DocumentContent


def extract_text_from_docx(file_path: str | Path) -> DocumentContent:
    """Extracts text, paragraphs, and font hints from a DOCX document.

    Args:
        file_path: Path to the .docx file.

    Returns:
        DocumentContent containing the extracted text, paragraph count, and font hints.

    Raises:
        FileNotFoundError: If the file does not exist.
        ValueError: If file is not a valid DOCX document.
    """
    path = Path(file_path)
    if not path.is_file():
        raise FileNotFoundError(f"DOCX file not found: {path}")

    try:
        doc = Document(str(path))
    except Exception as e:
        raise ValueError(f"Failed to parse DOCX file '{path}': {e}") from e

    paragraph_texts: list[str] = []
    font_hints: set[str] = set()

    for p in doc.paragraphs:
        if p.text:
            paragraph_texts.append(p.text)

        # Inspect run fonts (including complex script w:cs, ascii, hAnsi, eastAsia)
        for run in p.runs:
            if run.font and run.font.name:
                font_hints.add(run.font.name)
            try:
                rPr = run._r.find(qn("w:rPr"))
                if rPr is not None:
                    rFonts = rPr.find(qn("w:rFonts"))
                    if rFonts is not None:
                        for attr in ("w:cs", "w:ascii", "w:hAnsi", "w:eastAsia"):
                            val = rFonts.get(qn(attr))
                            if val:
                                font_hints.add(val)
            except Exception:
                pass

    # Also inspect tables if any
    for table in doc.tables:
        for row in table.rows:
            row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_texts:
                paragraph_texts.append(" | ".join(row_texts))

    full_text = "\n\n".join(paragraph_texts).strip()
    total_chars = len(full_text)
    font_hint = ", ".join(sorted(font_hints)) if font_hints else None

    return DocumentContent(
        text=full_text,
        format="docx",
        file_path=str(path.resolve()),
        font_hint=font_hint,
        page_count=max(1, (len(paragraph_texts) + 19) // 20),  # Heuristic page estimation
        character_count=total_chars,
        metadata={
            "paragraph_count": len(paragraph_texts),
            "font_hints": list(font_hints),
        },
    )
