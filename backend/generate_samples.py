"""Generates realistic Marathi sample documents for Phase 17 testing."""
from pathlib import Path
from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from pypdf import PdfWriter

SAMPLES_DIR = Path("samples")
SAMPLES_DIR.mkdir(exist_ok=True)

# 1. TXT sample (Mangal/UTF-8)
txt_content = (
    "महाराष्ट्र भारत देशातील २८ घटक राज्यांपैकी एक राज्य आहे.\n\n"
    "महाराष्ट्राची राजधानी मुंबई ही देशाची आर्थिक राजधानी मानली जाते. "
    "नागपूर ही महाराष्ट्राची उपराजधानी आहे.\n\n"
    "मराठी ही या राज्याची अधिकृत व सर्वाधिक बोलली जाणारी भाषा आहे. "
    "महाराष्ट्राला संतांची आणि शूरवीरांची भूमी म्हटले जाते. "
    "छत्रपती शिवाजी महाराजांनी येथे हिंदवी स्वराज्याची स्थापना केली.\n"
)
(SAMPLES_DIR / "sample_mangal.txt").write_text(txt_content, encoding="utf-8")
print("Generated samples/sample_mangal.txt")

# 2. DOCX sample (Kokila font hint)
docx = Document()
p1 = docx.add_paragraph()
r1 = p1.add_run("महाराष्ट्र राज्य माहिती व तंत्रज्ञान")
r1.font.name = "Kokila"
r1.bold = True

p2 = docx.add_paragraph()
r2 = p2.add_run(
    "माहिती तंत्रज्ञान क्षेत्रात महाराष्ट्राने उल्लेखनीय प्रगती केली आहे. "
    "पुणे व मुंबई ही प्रमुख माहिती तंत्रज्ञान केंद्रे म्हणून विकसित झाली आहेत. "
    "संगणक अभियांत्रिकी व कृत्रिम बुद्धिमत्ता या क्षेत्रात राज्यातील तरुणांचा मोठा सहभाग आहे."
)
r2.font.name = "Kokila"

docx.save(str(SAMPLES_DIR / "sample_kokila.docx"))
print("Generated samples/sample_kokila.docx")

# 3. PDF sample (Scanned / Image-only PDF simulation)
blank_writer = PdfWriter()
blank_writer.add_blank_page(width=595, height=842)
blank_writer.add_blank_page(width=595, height=842)
with open(SAMPLES_DIR / "sample_scanned.pdf", "wb") as f:
    blank_writer.write(f)
print("Generated samples/sample_scanned.pdf")
