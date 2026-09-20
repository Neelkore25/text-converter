"""Tests for PDF and DOCX Document Extraction and Output (Phases 13 & 14)."""
from pathlib import Path
import pytest
from docx import Document
from pypdf import PdfReader, PdfWriter

from backend.app.conversion.unicode_to_shivaji import convert_unicode_to_shivaji
from backend.app.extraction.docx_extractor import extract_text_from_docx
from backend.app.extraction.pdf_extractor import extract_text_from_pdf
from backend.app.output.docx_output import TARGET_FONT_NAME, write_converted_docx
from backend.app.output.pdf_output import write_converted_pdf


@pytest.fixture
def sample_docx(tmp_path: Path) -> Path:
    """Creates a sample Marathi DOCX document with Mangal font hint."""
    doc = Document()
    p1 = doc.add_paragraph()
    r1 = p1.add_run("महाराष्ट्र माझा देश आहे.")
    r1.font.name = "Mangal"

    p2 = doc.add_paragraph()
    r2 = p2.add_run("मला संगणक अभियांत्रिकी आवडते.")
    r2.font.name = "Mangal"

    p_path = tmp_path / "sample_input.docx"
    doc.save(str(p_path))
    return p_path


def test_docx_extraction_and_font_hint(sample_docx: Path):
    """Phase 14: Verify extraction of paragraphs and font hint from DOCX."""
    doc_content = extract_text_from_docx(sample_docx)
    assert doc_content.format == "docx"
    assert "महाराष्ट्र" in doc_content.text
    assert doc_content.font_hint is not None
    assert "Mangal" in doc_content.font_hint
    assert doc_content.character_count > 0


def test_docx_output_and_font_styling(tmp_path: Path):
    """Phase 14: Verify output DOCX has Shivaji01 Normal font explicitly set on runs."""
    converted_text = "maharaYT` maaJaa doSa Aaho.\n\nmalaa saMgaNak AiBayaaMi~kI AavaDto."
    out_path = tmp_path / "output_styled.docx"

    written = write_converted_docx(converted_text, out_path)
    assert written.exists()
    assert written.stat().st_size > 0

    # Re-open and verify font styling
    doc = Document(str(written))
    assert len(doc.paragraphs) == 2

    # Check font name on runs
    for p in doc.paragraphs:
        for r in p.runs:
            assert r.font.name == TARGET_FONT_NAME


def test_pdf_output_generation(tmp_path: Path):
    """Phase 13: Verify PDF generation embedding Shivaji01 font."""
    converted_text = "maharaYT` maaJaa doSa Aaho.\n\nmalaa saMgaNak AiBayaaMi~kI AavaDto."
    out_pdf = tmp_path / "output_sample.pdf"

    written_pdf = write_converted_pdf(converted_text, out_pdf)
    assert written_pdf.exists()
    assert written_pdf.stat().st_size > 0

    # Verify with pypdf
    reader = PdfReader(str(written_pdf))
    assert len(reader.pages) >= 1
    extracted = reader.pages[0].extract_text()
    assert len(extracted) > 0


def test_scanned_pdf_detection(tmp_path: Path):
    """Phase 13: Verify detection of scanned/image-only PDFs (no extractable text)."""
    # Create an empty 1-page PDF
    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)
    blank_pdf = tmp_path / "blank_scanned.pdf"
    with open(blank_pdf, "wb") as f:
        writer.write(f)

    doc_content = extract_text_from_pdf(blank_pdf)
    assert doc_content.is_scanned_image_only is True
    assert "OCR is required" in doc_content.metadata["notice"]


def test_full_docx_to_docx_roundtrip(sample_docx: Path, tmp_path: Path):
    """Full End-to-End: DOCX In -> Extract -> Convert -> DOCX Out."""
    extracted = extract_text_from_docx(sample_docx)
    res = convert_unicode_to_shivaji(extracted.text)
    assert res.warning_count == 0

    out_file = tmp_path / "final_converted.docx"
    written = write_converted_docx(res.converted_text, out_file)
    assert written.exists()

    reopened = extract_text_from_docx(written)
    assert reopened.character_count > 0
    assert "Mangal" not in (reopened.font_hint or "")
