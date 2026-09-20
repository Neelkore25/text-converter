"""Regression tests for Bug Audit & Fixes (Bugs 1 - 4).

Covers:
- BUG 1: Complex script font preservation (w:cs='Mangal', w:eastAsia, w:ascii, w:hAnsi)
- BUG 2: Paragraph structure / spacing preservation (5 paragraphs remain 5 paragraphs)
- BUG 3: HTML5 drag-and-drop dropzones and handlers
- BUG 4: History persistence, serialization with download_url, and retrieval
"""
from __future__ import annotations

from pathlib import Path
import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from fastapi.testclient import TestClient

from backend.app.api.server import app
from backend.app.extraction.docx_extractor import extract_text_from_docx
from backend.app.output.docx_output import set_run_font, write_converted_docx

client = TestClient(app)


def test_bug1_docx_complex_script_font_detection(tmp_path: Path):
    """BUG 1: Verify extract_text_from_docx detects Mangal when stored in w:rFonts w:cs."""
    docx_file = tmp_path / "mangal_complex_script.docx"
    doc = Document()
    p = doc.add_paragraph()
    run = p.add_run("महाराष्ट्र माझा देश आहे.")

    # Simulate Microsoft Word Devanagari complex-script font assignment
    rPr = run._r.get_or_add_rPr()
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:cs"), "Mangal")
    rFonts.set(qn("w:ascii"), "Calibri")  # Default ascii slot in Word
    rPr.append(rFonts)

    doc.save(str(docx_file))

    extracted = extract_text_from_docx(docx_file)
    assert extracted.font_hint is not None
    assert "Mangal" in extracted.font_hint
    assert "Mangal" in extracted.metadata["font_hints"]


def test_bug1_docx_output_sets_all_four_font_slots(tmp_path: Path):
    """BUG 1: Verify set_run_font explicitly sets ascii, hAnsi, eastAsia, AND cs font slots."""
    docx_file = tmp_path / "test_out_font_slots.docx"
    doc = Document()
    p = doc.add_paragraph()
    run = p.add_run("maharaYT`")
    set_run_font(run, "Shivaji01 Normal")
    doc.save(str(docx_file))

    # Re-open and verify XML attributes
    saved_doc = Document(str(docx_file))
    rPr = saved_doc.paragraphs[0].runs[0]._r.find(qn("w:rPr"))
    assert rPr is not None
    rFonts = rPr.find(qn("w:rFonts"))
    assert rFonts is not None

    assert rFonts.get(qn("w:ascii")) == "Shivaji01 Normal"
    assert rFonts.get(qn("w:hAnsi")) == "Shivaji01 Normal"
    assert rFonts.get(qn("w:eastAsia")) == "Shivaji01 Normal"
    assert rFonts.get(qn("w:cs")) == "Shivaji01 Normal"


def test_bug2_paragraph_structure_preservation_5_paragraphs(tmp_path: Path):
    """BUG 2: Verify 5 separate paragraphs remain 5 distinct paragraphs end-to-end."""
    sample_paragraphs = [
        "परिच्छेद क्रमांक १: महाराष्ट्र भारत देशातील एक प्रमुख राज्य आहे.",
        "परिच्छेद क्रमांक २: मुंबई ही महाराष्ट्राची राजधानी आहे.",
        "परिच्छेद क्रमांक ३: मराठी ही महाराष्ट्राची अधिकृत राज्यभाषा आहे.",
        "परिच्छेद क्रमांक ४: सह्याद्री पर्वतरांग महाराष्ट्राच्या पश्चिमेस आहे.",
        "परिच्छेद क्रमांक ५: संतांची भूमी म्हणून महाराष्ट्राची ख्याती आहे."
    ]

    # 1. Create a 5-paragraph DOCX
    src_docx = tmp_path / "five_paragraphs.docx"
    doc = Document()
    for text in sample_paragraphs:
        doc.add_paragraph(text)
    doc.save(str(src_docx))

    # 2. Extract and check paragraph count
    extracted = extract_text_from_docx(src_docx)
    assert extracted.metadata["paragraph_count"] == 5
    extracted_paras = extracted.text.split("\n\n")
    assert len(extracted_paras) == 5

    # 3. Export to DOCX and check generated paragraphs
    out_docx = tmp_path / "five_paragraphs_out.docx"
    write_converted_docx(extracted.text, out_docx)

    out_doc = Document(str(out_docx))
    assert len(out_doc.paragraphs) == 5
    for idx, p in enumerate(out_doc.paragraphs):
        assert len(p.text.strip()) > 0
        assert p.runs[0].font.name == "Shivaji01 Normal"


def test_bug3_html_drag_and_drop_elements_and_handlers():
    """BUG 3: Verify index.html contains dropzones, preventDefault wiring, and dragover styling."""
    html_path = Path("index.html")
    assert html_path.exists()
    content = html_path.read_text(encoding="utf-8")

    # Drop zone IDs
    assert 'id="drop-zone"' in content
    assert 'id="quick-drop-zone"' in content

    # CSS dragover feedback
    assert ".upload-box.dragover" in content
    assert "accent-cyan" in content

    # JS event listeners with preventDefault
    assert "setupDragAndDrop" in content
    assert "e.preventDefault()" in content
    assert "e.stopPropagation()" in content
    assert "dataTransfer.files" in content


def test_bug4_history_api_serialization_and_download_links():
    """BUG 4: Verify /api/history returns valid download_url, output_filename, and records conversions."""
    # 1. Convert a file via API
    marathi_text = "महाराष्ट्र हे भारतातील एक महान राज्य आहे."
    file_bytes = marathi_text.encode("utf-8")
    files = {"file": ("audit_history_test.txt", file_bytes, "text/plain")}
    data = {"source_font": "auto", "output_format": "docx"}

    post_resp = client.post("/api/convert/file", files=files, data=data)
    assert post_resp.status_code == 200
    post_data = post_resp.json()
    assert "output_filename" in post_data
    assert "download_url" in post_data

    # 2. Query /api/history
    hist_resp = client.get("/api/history?limit=10")
    assert hist_resp.status_code == 200
    records = hist_resp.json()
    assert len(records) > 0

    latest = records[0]
    assert "output_filename" in latest
    assert "download_url" in latest
    assert latest["output_filename"].startswith("converted_audit_history_test.docx")
    assert latest["download_url"].startswith("/api/download/")
    assert latest["status"] is not None

    # 3. Verify download works through the URL provided in history
    dl_resp = client.get(latest["download_url"])
    assert dl_resp.status_code == 200
    assert len(dl_resp.content) > 0


def test_end_to_end_mangal_docx_paste_simulation(tmp_path: Path):
    """End-to-end: Simulate pasting a 5-paragraph Mangal DOCX with complex script font hints."""
    # Create source Word document with Mangal complex script
    doc = Document()
    paragraphs = [
        "पहिला परिच्छेद: छत्रपती शिवाजी महाराज हे एक महान राजे होते.",
        "दुसरा परिच्छेद: त्यांनी स्वराज्य स्थापन केले.",
        "तिसरा परिच्छेद: गड-किल्ले जिंकून त्यांनी रयतेचे राज्य निर्माण केले.",
        "चौथा परिच्छेद: अष्टप्रधान मंडळ स्थापन करून उत्तम प्रशासन दिले.",
        "पाचवा परिच्छेद: आजही त्यांचे विचार संपूर्ण देशाला प्रेरणा देतात."
    ]
    for p_text in paragraphs:
        p = doc.add_paragraph()
        run = p.add_run(p_text)
        rPr = run._r.get_or_add_rPr()
        rFonts = OxmlElement("w:rFonts")
        rFonts.set(qn("w:cs"), "Mangal")
        rFonts.set(qn("w:ascii"), "Calibri")
        rPr.append(rFonts)

    src_docx = tmp_path / "real_mangal_input.docx"
    doc.save(str(src_docx))

    # Send to conversion endpoint
    with open(src_docx, "rb") as f:
        files = {"file": ("real_mangal_input.docx", f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        data = {"source_font": "auto", "output_format": "docx"}
        resp = client.post("/api/convert/file", files=files, data=data)

    assert resp.status_code == 200
    res_json = resp.json()
    assert res_json["source_font"] == "Mangal"
    assert res_json["target_font"] == "Shivaji01 Normal"
    assert res_json["warning_count"] == 0

    # Download output DOCX and verify
    dl_resp = client.get(res_json["download_url"])
    assert dl_resp.status_code == 200

    out_file = tmp_path / "verified_output.docx"
    out_file.write_bytes(dl_resp.content)

    out_doc = Document(str(out_file))
    assert len(out_doc.paragraphs) == 5

    for p in out_doc.paragraphs:
        assert len(p.runs) > 0
        rFonts = p.runs[0]._r.find(qn("w:rPr")).find(qn("w:rFonts"))
        assert rFonts.get(qn("w:cs")) == "Shivaji01 Normal"
        assert rFonts.get(qn("w:ascii")) == "Shivaji01 Normal"
        assert rFonts.get(qn("w:hAnsi")) == "Shivaji01 Normal"
        assert rFonts.get(qn("w:eastAsia")) == "Shivaji01 Normal"


def test_docx_with_table_header_footer_formatting_preserved(tmp_path: Path):
    """Verify DOCX with tables, headers, and footers has all text converted and formatting preserved."""
    from backend.app.output.docx_output import transform_existing_docx

    src_docx = tmp_path / "table_and_header.docx"
    doc = Document()

    # 1. Add Header
    section = doc.sections[0]
    header = section.header
    header_p = header.paragraphs[0]
    header_run = header_p.add_run("दस्तावेज शीर्षक: महाराष्ट्र परिचय")
    set_run_font(header_run, "Mangal")

    # 2. Add Footer
    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_run = footer_p.add_run("पृष्ठ १: अधिकृत अहवाल")
    set_run_font(footer_run, "Mangal")

    # 3. Add Body Paragraph with Bold/Italic runs
    p = doc.add_paragraph()
    r1 = p.add_run("ठळक मजकूर: ")
    r1.bold = True
    set_run_font(r1, "Mangal")
    r2 = p.add_run("तिरपा मजकूर: महाराष्ट्र.")
    r2.italic = True
    set_run_font(r2, "Mangal")

    # 4. Add Table with cells
    table = doc.add_table(rows=2, cols=2)
    cell_texts = [
        ["जिल्हा", "मुख्यालय"],
        ["पुणे", "पुणे शहर"]
    ]
    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            cell_p = cell.paragraphs[0]
            c_run = cell_p.add_run(cell_texts[r_idx][c_idx])
            set_run_font(c_run, "Mangal")

    doc.save(str(src_docx))

    # Transform document
    out_docx = tmp_path / "table_and_header_converted.docx"
    transform_existing_docx(src_docx, out_docx)

    # Verify converted output
    res_doc = Document(str(out_docx))

    # Verify Header
    res_hdr_run = res_doc.sections[0].header.paragraphs[0].runs[0]
    hdr_fonts = res_hdr_run._r.find(qn("w:rPr")).find(qn("w:rFonts"))
    assert hdr_fonts.get(qn("w:cs")) == "Shivaji01 Normal"
    assert hdr_fonts.get(qn("w:ascii")) == "Shivaji01 Normal"

    # Verify Footer
    res_ftr_run = res_doc.sections[0].footer.paragraphs[0].runs[0]
    ftr_fonts = res_ftr_run._r.find(qn("w:rPr")).find(qn("w:rFonts"))
    assert ftr_fonts.get(qn("w:cs")) == "Shivaji01 Normal"

    # Verify Bold/Italic preserved
    res_p = res_doc.paragraphs[0]
    assert res_p.runs[0].bold is True
    assert res_p.runs[1].italic is True

    # Verify Table Cells
    res_table = res_doc.tables[0]
    assert len(res_table.rows) == 2
    for row in res_table.rows:
        for cell in row.cells:
            cell_run = cell.paragraphs[0].runs[0]
            c_fonts = cell_run._r.find(qn("w:rPr")).find(qn("w:rFonts"))
            assert c_fonts.get(qn("w:cs")) == "Shivaji01 Normal"
            assert c_fonts.get(qn("w:ascii")) == "Shivaji01 Normal"
            assert c_fonts.get(qn("w:hAnsi")) == "Shivaji01 Normal"
            assert c_fonts.get(qn("w:eastAsia")) == "Shivaji01 Normal"


def test_legacy_fonts_presence_and_mapping_status():
    """Verify font file presence and mapping table presence are tracked separately for legacy fonts."""
    from backend.app.fonts.source_fonts import get_source_font_by_id

    for font_id in ("kruti_dev", "devlys", "kantal"):
        font = get_source_font_by_id(font_id)
        assert font is not None
        # File is not present in assets/fonts/source
        assert font.font_file_available is False
        # Mapping table is not present
        assert font.mapping_table_available is False
        assert "requires font" in font.status.value.lower()

    # Verify API reports both attributes
    resp = client.get("/api/fonts")
    assert resp.status_code == 200
    data = resp.json()
    for sf in data["source_fonts"]:
        if sf["font_type"] == "legacy_encoded":
            assert sf["font_file_available"] is False
            assert sf["mapping_table_available"] is False


def test_client_side_jszip_docx_pipeline_markup():
    """Verify index.html contains JSZip, magic-byte inspection, formatting preservation, and detected font badge."""
    html_content = Path("index.html").read_text(encoding="utf-8")

    # JSZip library inclusion
    assert "jszip.min.js" in html_content

    # Magic bytes check for ZIP (0x50 0x4B) and legacy .doc (0xD0 0xCF 0x11 0xE0)
    assert "0x50 && headerBytes[1] === 0x4B" in html_content
    assert "0xD0 && headerBytes[1] === 0xCF" in html_content
    assert "Legacy .doc format is not supported" in html_content

    # word/document.xml extraction & parsing
    assert 'zip.file("word/document.xml")' in html_content
    assert "DOMParser" in html_content

    # Font inspection and display
    assert 'id="topbar-source-font"' in html_content
    assert 'id="detected-source-badge"' in html_content
    assert 'id="detected-font-name"' in html_content

    # Run-level font assignment across all 4 slots
    assert 'rFonts.setAttribute(\'w:ascii\', \'Shivaji01 Normal\')' in html_content
    assert 'rFonts.setAttribute(\'w:cs\', \'Shivaji01 Normal\')' in html_content
    assert 'rFonts.setAttribute(\'w:eastAsia\', \'Shivaji01 Normal\')' in html_content

    # Re-zipping and Blob generation
    assert "generateAsync" in html_content
    assert "application/vnd.openxmlformats-officedocument.wordprocessingml.document" in html_content
